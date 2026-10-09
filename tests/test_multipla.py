import copy
import json
from pathlib import Path

import openpyxl

from exportador_multipla import gerar_origem, ler_multipla, validar

BUEIRO = json.loads((Path(__file__).parent / "bueiro.json").read_text(encoding="utf-8"))


def _sem_custo(orc):
    # O custo unitário sai como fórmula =AG (igual à PM) e só tem valor depois que o Excel recalcula.
    orc = copy.deepcopy(orc)
    for item in orc["itens"]:
        item.pop("custo_unitario", None)
    return orc


def test_ida_e_volta_bueiro(tmp_path):
    saida = gerar_origem(BUEIRO, tmp_path / "origem.xlsx")
    assert _sem_custo(ler_multipla(saida)) == _sem_custo(BUEIRO)


def test_custo_unitario_vai_em_ag_com_formula_em_u(tmp_path):
    saida = gerar_origem(BUEIRO, tmp_path / "origem.xlsx")
    ws = openpyxl.load_workbook(saida)["ORÇAMENTO"]
    assert ws["U18"].value == "=AG18"
    assert ws["AG18"].value == 15.66


def test_orcamento_maior_ajusta_ultima_linha(tmp_path):
    orc = copy.deepcopy(BUEIRO)
    servicos = [i for i in orc["itens"] if i["nivel"] == "Serviço"]
    orc["itens"] += copy.deepcopy(servicos) * 3
    saida = gerar_origem(orc, tmp_path / "origem.xlsx")

    wb = openpyxl.load_workbook(saida)
    ultima = 15 + len(orc["itens"]) + 1
    assert wb.defined_names["ORÇAMENTO.lastrow"].attr_text == f"ORÇAMENTO!${ultima}:${ultima}"
    assert wb.defined_names["Import.Nível"].attr_text.endswith(f"$M${ultima},-1,0)")
    assert len(ler_multipla(saida)["itens"]) == len(orc["itens"])


def test_marcadores_exigidos_pela_macro_da_caixa(tmp_path):
    wb = openpyxl.load_workbook(gerar_origem(BUEIRO, tmp_path / "origem.xlsx"))
    assert wb["MENU"]["J1"].value == "PM"
    for nome in ["Objeto", "ORÇAMENTO.firstrow", "ORÇAMENTO.LinhaPadrão", "CÁLCULO.firstcol",
                 "CÁLCULO.lastcol", "CRONO.NivelExibicao", "CRONO.Parcela1", "PLE.Medicao", "BM.medicao"]:
        assert nome in wb.defined_names


def test_validar_aponta_erros():
    orc = copy.deepcopy(BUEIRO)
    orc["itens"][2]["fonte"] = "XYZ"
    orc["itens"][3]["codigo"] = ""
    erros = validar(orc)
    assert any("fonte inválida" in e for e in erros)
    assert any("sem código" in e for e in erros)
