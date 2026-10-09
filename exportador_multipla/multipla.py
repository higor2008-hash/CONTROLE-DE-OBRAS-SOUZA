"""Leitura e geração de arquivos compatíveis com a Planilha Múltipla (PM) da Caixa.

A PM tem um botão "Importar de outra Planilha Múltipla" (macro Importar_PM3) que
lê, no arquivo de origem, os nomes definidos que começam com "Import." e grava os
valores nas células editáveis da PM de destino, inserindo as linhas necessárias.

Em vez de editar o .xlsm oficial (macros, botões e proteção), geramos um arquivo
de ORIGEM simples (.xlsx) com as mesmas abas e nomes definidos. O usuário abre a PM
oficial em branco e importa esse arquivo: a própria macro da Caixa faz o resto.
"""
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from .names import resolve

HERE = Path(__file__).parent
TEMPLATE = json.loads((HERE / "pm316_names.json").read_text(encoding="utf-8"))

# Linhas da PM v3.16
ORC_FIRST = 15          # linha do LOTE (ORÇAMENTO.firstrow / CÁLCULO.firstrow)
TEMPLATE_LAST = 22      # ORÇAMENTO.lastrow no arquivo de referência
N_FRENTES = 10          # colunas Q..Z da aba CÁLCULO

DADOS_CAMPOS = {
    "recurso": "import.recurso",
    "proponente": "Import.Proponente",
    "municipio": "Import.Município",
    "operacao": "Import.CR",
    "transferegov": "Import.TransfereGOV",
    "repasse": "Import.Repasse",
    "contrapartida": "Import.Contrapartida",
    "apelido": "Import.Apelido",
    "descricao_lote": "Import.DescLote",
    "desoneracao": "Import.Desoneracao",
    "data_base": "Import.DataBase",
    "data_preenchimento": "Import.DataPreenchimento",
    "arredondamento": "Import.TipoArredondamento",
    "regime_execucao": "Import.RegimeExecução",
}
LISTAS_DADOS = {
    "responsavel_orcamento": "Import.RespOrçamento",   # [nome, crea/cau, art/rrt]
    "prefeito": "IMport.Prefeito",                      # [nome, cargo]
}
ITEM_COLUNAS = {
    "nivel": "Import.Nível",
    "fonte": "Import.Fonte",
    "codigo": "Import.Código",
    "descricao": "Import.Descrição",
    "unidade": "Import.Unidade",
    "custo_unitario": "Import.CustoUnitário",
    "bdi": "Import.OpcaoBDI",
    "recurso": "Import.ORÇAMENTO.DivRecurso",
    "memoria": "Import.PLQ.MemCalc",
}
NIVEIS = {"Meta", "Nível 2", "Nível 3", "Nível 4", "Serviço"}
FONTES = {"SINAPI", "SINAPI-I", "SICRO", "Composição", "Cotação"}


def _jsonable(v):
    return v.isoformat()[:10] if isinstance(v, datetime) else v


def _read(wb, name):
    sh, c1, r1, c2, r2 = resolve(wb.defined_names[name].attr_text)
    ws = wb[sh]
    return [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


def ler_multipla(caminho):
    """Lê uma PM (ou um arquivo de origem gerado aqui) e devolve o orçamento como dict."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    col = lambda name: [row[0] for row in _read(wb, name)]

    orc = {"dados": {}, "bdi": {}, "itens": []}
    for campo, nome in DADOS_CAMPOS.items():
        orc["dados"][campo] = _jsonable(_read(wb, nome)[0][0])
    for campo, nome in LISTAS_DADOS.items():
        orc["dados"][campo] = [r[0] for r in _read(wb, nome)]

    orc["bdi"] = {
        "tipo": _read(wb, "Import.BDI.Tipo1")[0][0],
        "parcelas": dict(zip(["AC", "SG", "R", "DF", "L", "CP"], col("Import.BDI.Det1"))),
        "iss_base": _read(wb, "Import.BDI.ISS")[0][0],
        "iss_aliquota": _read(wb, "Import.BDI.ISS")[1][0],
    }

    colunas = {campo: col(nome) for campo, nome in ITEM_COLUNAS.items()}
    frentes = _read(wb, "Import.PLQ")
    for i, nivel in enumerate(colunas["nivel"]):
        if not nivel:
            continue
        item = {campo: _jsonable(colunas[campo][i]) for campo in ITEM_COLUNAS}
        if isinstance(item["codigo"], (int, float)):
            item["codigo"] = str(int(item["codigo"]))
        if nivel == "Serviço":
            q = [v or 0 for v in frentes[i]]
            while q and q[-1] == 0:
                q.pop()
            item["quantidade_frentes"] = q
        else:
            for k in ("fonte", "codigo", "unidade", "custo_unitario", "memoria"):
                item.pop(k)
        orc["itens"].append(item)
    return orc


def validar(orc):
    erros = []
    itens = orc["itens"]
    if not itens or itens[0]["nivel"] != "Meta":
        erros.append("O primeiro item precisa ser uma Meta.")
    for n, it in enumerate(itens, 1):
        if it["nivel"] not in NIVEIS:
            erros.append(f"Item {n}: nível inválido {it['nivel']!r}.")
        if it["nivel"] == "Serviço":
            if it.get("fonte") not in FONTES:
                erros.append(f"Item {n}: fonte inválida {it.get('fonte')!r}.")
            if not it.get("codigo"):
                erros.append(f"Item {n}: serviço sem código.")
            if len(it.get("quantidade_frentes", [])) > N_FRENTES:
                erros.append(f"Item {n}: mais de {N_FRENTES} frentes de obra.")
    return erros


def _nome_ajustado(texto, ultima):
    """Ajusta referências à última linha do ORÇAMENTO/CÁLCULO para o tamanho do orçamento."""
    texto = re.sub(r"((?:ORÇAMENTO|CÁLCULO)!\$[A-Z]+\$)22\b", rf"\g<1>{ultima}", texto)
    texto = re.sub(r"((?:ORÇAMENTO|CÁLCULO)!\$[A-Z]+\$\d+:\$[A-Z]+\$)22\b", rf"\g<1>{ultima}", texto)
    return re.sub(r"((?:ORÇAMENTO|CÁLCULO)!\$)22:\$22\b", rf"\g<1>{ultima}:${ultima}", texto)


def gerar_origem(orc, saida):
    """Gera o arquivo de origem (.xlsx) para importar na PM oficial."""
    erros = validar(orc)
    if erros:
        raise ValueError("\n".join(erros))

    n = len(orc["itens"])
    ultima = ORC_FIRST + n + 1
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome in TEMPLATE["sheets"]:
        wb.create_sheet(nome)

    nomes = {}
    for nome, texto in TEMPLATE["names"].items():
        if nome.startswith("_xlnm"):
            continue
        texto = _nome_ajustado(texto, ultima)
        nomes[nome] = texto
        wb.defined_names[nome] = DefinedName(nome, attr_text=texto)

    def put(nome, valores):
        sh, c1, r1, c2, r2 = resolve(nomes[nome])
        ws = wb[sh]
        for i, linha in enumerate(valores):
            for j, v in enumerate(linha):
                if v not in (None, ""):
                    ws.cell(r1 + i, c1 + j, v)

    # Marcadores que a macro da Caixa confere na origem
    wb["MENU"]["J1"] = "PM"
    wb["MENU"]["J2"] = TEMPLATE["versao"]
    wb["MENU"]["O3"] = 1           # Tipo de orçamento: Proposto
    wb["MENU"]["O4"] = 1           # Acompanhamento: PLE
    wb["CRONO"]["H10"] = 2         # nível de exibição do cronograma
    wb["CÁLCULO"]["K12"] = "Nº AGRUPADOR DE EVENTOS"
    wb["ORÇAMENTO"].cell(ORC_FIRST, 13, "LOTE")
    wb["ORÇAMENTO"].cell(ultima, 1, -1)

    d = orc["dados"]
    for campo, nome in DADOS_CAMPOS.items():
        v = d.get(campo)
        if campo in ("data_base", "data_preenchimento") and isinstance(v, str):
            v = datetime.fromisoformat(v)
        put(nome, [[v]])
    for campo, nome in LISTAS_DADOS.items():
        put(nome, [[v] for v in d.get(campo, [])])

    b = orc["bdi"]
    put("Import.BDI.Tipo1", [[b["tipo"]]])
    put("Import.BDI.Det1", [[b["parcelas"].get(k)] for k in ["AC", "SG", "R", "DF", "L", "CP"]])
    put("Import.BDI.ISS", [[b.get("iss_base")], [b.get("iss_aliquota")]])
    put("Import.BDI.Tipo2", [["(SELECIONAR)"]])
    put("Import.BDI.Tipo3", [["(SELECIONAR)"]])

    for campo, nome in ITEM_COLUNAS.items():
        put(nome, [[it.get(campo)] for it in orc["itens"]])
    put("Import.PLQ", [it.get("quantidade_frentes", []) for it in orc["itens"]])
    put("Import.FrenteDeObra", [[1] + [None] * (N_FRENTES - 1)])

    # Custo unitário como fórmula =AG (como na PM), com o custo de referência em AG,
    # para funcionar com e sem a opção "Importar Fórmulas para o Custo Unitário".
    ws = wb["ORÇAMENTO"]
    for i, it in enumerate(orc["itens"]):
        r = ORC_FIRST + 1 + i
        if it["nivel"] == "Serviço":
            ws.cell(r, 33, it.get("custo_unitario") or 0)      # AG
            ws.cell(r, 21, f"=AG{r}")                           # U

    put("Import.QCI.Divisao", [["Proporcional"]] * 10)
    put("Import.Eventos.Nível", [["Automática, conforme os agrupadores Nível 2 do Orçamento"]])
    put("Import.POArred", [[True]] * 5)
    put("Import.BMArred", [[True]])

    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
    return saida


def main(argv):
    """python -m exportador_multipla ler ARQUIVO.xlsm | gerar ORCAMENTO.json SAIDA.xlsx
    | preencher ORCAMENTO.json MULTIPLA_MODELO.xlsm SAIDA.xlsm"""
    cmd = argv[1]
    if cmd == "ler":
        print(json.dumps(ler_multipla(argv[2]), ensure_ascii=False, indent=1))
    elif cmd == "preencher":
        from .preencher import preencher_multipla
        orc = json.loads(Path(argv[2]).read_text(encoding="utf-8"))
        print(preencher_multipla(orc, argv[3], argv[4]))
    elif cmd == "gerar":
        orc = json.loads(Path(argv[2]).read_text(encoding="utf-8"))
        print(gerar_origem(orc, argv[3]))
